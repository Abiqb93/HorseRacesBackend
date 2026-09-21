#!/usr/bin/env python3
"""Build breeding/fixtures/wathnan-2027.json from the consultant's 2027 matings pack.

Sources (all under --scratch unless noted):
  master_text.txt                 page-split text ("===== PAGE n (chars=...) =====") of the
                                  "Wathnan Matings Master File for 2027 Season" PDF, one report per mare
  INDEX_ROWS (below)              transcription of "Matings suggestions index for 2027" (PDF)
  zip/Wathnan summary.xlsx        sheets 'Mares' (summary grid) and 'Fillies in training'
  fixture/docx/, zip/             per-mare "<MARE> 27.docx" from the OneDrive working folder.  Missing ones
                                  are fetched from the 1.5 GB zip on Google Drive by HTTP range request
                                  (central directory from ziptail.bin), one request per second.
  ../breeding/data/horses-<yob>.csv.gz   Timeform export (repo) - only used for the foaling country of
                                  mares the index leaves unsuffixed, and for LADY VIVIAN (not in the index).

Usage:
  PYTHONPATH=<scratch>/pylib python3 scripts/build-wathnan-fixture.py [--scratch DIR] [--out FILE] [--no-fetch]

Only the stdlib plus openpyxl (for the xlsx) are needed.
"""
import argparse
import csv
import datetime as dt
import gzip
import json
import os
import re
import struct
import subprocess
import sys
import time
import zipfile
import xml.etree.ElementTree as ET

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
DEFAULT_SCRATCH = "/tmp/claude-0/-home-user/ba406b5d-7269-5dcb-b35d-3cd8ae3b88af/scratchpad"
DEFAULT_OUT = os.path.join(REPO, "breeding", "fixtures", "wathnan-2027.json")
TIMEFORM_DIR = os.path.join(REPO, "breeding", "data")

ZIP_URL = ("https://drive.usercontent.google.com/download?id=1FeoR7z15WCI95k3TctEn7IcTntto6N2D"
           "&export=download&confirm=t")
ZIP_SIZE = 1561030430
ZIP_TAIL_BYTES = 1500000
ZIP_PACK_PREFIX = "WATHNAN MATINGS 2026/"

SEASON = 2027
PREV_SEASON = "2026"

# ----------------------------------------------------------------------------------------------------
# Band index, transcribed from "Matings suggestions index for 2027" (PDF).
# Columns: MARE (country in brackets = where FOALED) | YOB | SIRE | DAM | CURRENT LOCATION | COUNTRY (where
# she STANDS) | Foaled 2026 | 2026 MATING PLAN | STATUS | LSD | 2027 MATING SUGGESTION.  "—" = blank cell.
# ----------------------------------------------------------------------------------------------------
INDEX_TEXT = """
CIRCIOS (IRE) | 2022 | KINGMAN | COULD IT BE LOVE | Newsells | UK | — | Frankel | In foal | 24 Feb 26 | Not This Time/ Night of Thunder
CRIMSON ADVOCATE (USA) | 2021 | NYQUIST | CITIZEN ADVOCATE | Whatton Manor | UK | — | Frankel | In foal | 26 Feb 26 | Not This Time/ Night of Thunder
DARE TO DREAM (FR) | 2021 | CAMELOT | DEBUTANTE | Kellsgrange Stud | UK | — | Sea The Stars | In foal | 25 Mar 2026 | Night of Thunder / Lope de Vega
DRAMATISED (IRE) | 2020 | SHOWCASING | KATIE'S DIAMOND | Newsells | UK | 07.02.26 | Frankel | In foal | 7 Mar 2026 | Blue Point or No Nay Never
EXXTRA (FR) | 2020 | STARSPANGLEDBANNER | ROLLING STONE | Whatton Manor | UK | — | Frankel | In foal | 1 Mar 26 | Dubawi or Kingman
FALLEN ANGEL | 2021 | TOO DARN HOT | AGNES STEWART | — | — | Maiden | — | — | — | Frankel
FLORA OF BERMUDA | 2021 | DARK ANGEL | DUBAI POWER | — | — | Maiden | — | — | — | Night Of Thunder or Kingman
HEREDIA (GB) | 2019 | DARK ANGEL | NAKUTI | Lane's End | USA | — | Justify | In foal | 20 May 26 | Not This Time or Frankel
IMMENSITUDE (FR) | 2020 | LAWMAN | MA PETITE POULE | Lane's End | USA | — | Justify | In foal | 4 Mar 2026 | Not This Time / Night Of Thunder
INFINITE COSMOS (IRE) | 2020 | SEA THE STARS | WAILA | Lane's End | USA | 25.02.26 | Justify | In foal | 14 May 2026 | Night Of Thunder or Dubawi
JASNA'S SECRET (FR) | 2021 | GALIWAY | ORPENA | Whatton Manor | UK | — | Kingman | In foal | 24 Mar 26 | Justify / Lope de Vega
JULICA (FR) | 2021 | KENDARGENT | MISS RAIL LINK | Whatton Manor | UK | — | Frankel | In foal | 17 Feb 26 | Night of Thunder or Justify
LEOVANNI | 2021 | KODI BEAR | KASSANDRA | — | — | Maiden | — | — | — | Frankel
LOPE DE LILAS | 2021 | LOPE DE VEGA | GOLDEN GAZELLE | — | — | Maiden | — | — | — | Kingman or Too Darn Hot
MELO MELO (GB) | 2019 | GLENEAGLES | YOU LOOK SO GOOD | Whatton Manor | UK | 08.02.26 | Kingman | In foal | 16 Mar 2026 | Justify or Night Of Thunder
OCALA (GB) | 2015 | NATHANIEL | NIGHT CARNATION | Far Westfield | UK | 30.01.26 | Starspangledbanner | In foal | 27 Feb 26 | St Mark's Basilica or New Bay
OLENTIA (AUS) | 2019 | ZOUSTAR | MABKHARA | Whatton Manor | UK | — | Frankel | In foal | 16 Feb 26 | Justify or Frankel
ONE LOOK | 2021 | GLENEAGLES | HOLY SALT | — | — | Maiden | — | — | — | Sea The Stars or Justify
PEARLA | 2022 | SEA THE STARS | TRAFFIC JAM | — | — | Maiden | — | — | — | Lope de Vega or Frankel
REMARQUEE (GB) | 2020 | KINGMAN | REGARDEZ | Coolmore | Ire | 07.02.26 | Night Of Thunder | In foal | 12 Mar 2026 | Frankel or Justify
SCENIC (FR) | 2020 | LOPE DE VEGA | GHALYAH | Lane's End | USA | — | Justify | In foal | 24 Feb 2026 | Dubawi or Night Of Thunder
SERENE SERAPH (IRE) | 2021 | BLUE POINT | PACIFIC ANGEL | Lane's End | USA | — | Justify | In foal | 15 Mar 2026 | Sea The Stars or Frankel
SHINE ON | 2023 | HAVANA GREY | MOTSI MA BOATI | — | — | Maiden | — | — | — | Blue Point or No Nay Never
SOPRANO (IRE) | 2021 | STARSPANGLEDBANNER | LEALAS DAUGHTER | Kellsgrange Stud | UK | — | Night Of Thunder | In foal | 18 Feb 26 | Frankel or Night of Thunder
STAY ALERT (GB) | 2019 | FASTNET ROCK | STARFALA | Lane's End | USA | 03.03.26 | Justify | Not in foal | 18 Apr 2026 | Not This Time / Sea The Stars
SUMO SAM (GB) | 2020 | NATHANIEL | SEADUCED | Newsells | UK | 11.02.26 | Dubawi | In foal | 19 Apr 2026 | Kingman or Justify/NTT
"""
# The index writes "SHINE ON"; the master file, docx pack and Timeform all say SHINE ON ME.
INDEX_NAME_FIXES = {"SHINE ON": "SHINE ON ME"}

# LADY VIVIAN is in the working folder and the summary grid but not in the index or the master file.
DOCX_ONLY_MARES = ["LADY VIVIAN"]

STALLION_ALIASES = {"NTT": "NOT THIS TIME", "NOT": "NIGHT OF THUNDER", "STS": "SEA THE STARS",
                    "SSB": "STARSPANGLEDBANNER", "LDV": "LOPE DE VEGA"}
REGION_ALIASES = {"USA": "USA", "US": "USA", "AMERICA": "USA", "EU": "EU", "EUROPE": "EU", "UK": "UK"}
STANDS_IN = {"UK": "UK", "GB": "UK", "USA": "USA", "US": "USA", "IRE": "IRE", "IRELAND": "IRE"}

# ----------------------------------------------------------------------------------------------------
# text helpers
# ----------------------------------------------------------------------------------------------------
LIGATURES = {"ﬀ": "ff", "ﬁ": "fi", "ﬂ": "fl", "ﬃ": "ffi", "ﬄ": "ffl",
             "ﬅ": "st", "ﬆ": "st"}
COUNTRY_SUFFIX_RE = re.compile(r"\s*\((?:GB|IRE|USA|FR|GER|AUS|NZ|JPN|CAN|ARG|BRZ|CHI|SAF|ITY|SPA|UAE|KOR|"
                               r"TUR|IND|HK|SWE|NOR|DEN|POL|CZE|HUN|RUS|URU|PER|VEN|MEX|SLO|SIN|MOR)\)\s*$",
                               re.I)
MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep",
                                      "oct", "nov", "dec"], 1)}


def fix_ligatures(s):
    for k, v in LIGATURES.items():
        s = s.replace(k, v)
    return s


def norm_ws(s):
    return re.sub(r"\s+", " ", s or "").strip()


def horse_name(s):
    """UPPER-CASE horse name, straight apostrophes, no country suffix."""
    s = fix_ligatures(s or "").replace("’", "'").replace("‘", "'").replace("`", "'")
    s = norm_ws(s)
    s = COUNTRY_SUFFIX_RE.sub("", s)
    s = s.upper().strip(" .")
    return STALLION_ALIASES.get(s, s)


def blank(v):
    return v is None or str(v).strip() in ("", "—", "-", "–")


def iso_date(s):
    """'24 Feb 26' | '25 Mar 2026' | '07.02.26' | '7/03/26' -> 'YYYY-MM-DD' (2-digit years are 20xx)."""
    if blank(s):
        return None
    s = str(s).strip()
    m = re.match(r"^(\d{1,2})[./](\d{1,2})[./](\d{2,4})$", s)
    if m:
        d, mo, y = int(m[1]), int(m[2]), int(m[3])
    else:
        m = re.match(r"^(\d{1,2})(?:st|nd|rd|th)?\s+([A-Za-z]{3})[a-z]*\.?\s+(\d{2,4})$", s)
        if not m:
            raise ValueError(f"unparsed date {s!r}")
        d, mo, y = int(m[1]), MONTHS[m[2].lower()], int(m[3])
    if y < 100:
        y += 2000
    return f"{y:04d}-{mo:02d}-{d:02d}"


def join_wrapped_lines(lines):
    """Join PDF-wrapped lines into one whitespace-normalised string.  A line ending in a hyphen that is
    glued to a word ("half-" / "sister") is joined without a space; "Derby - " (spaced dash) is not."""
    out = ""
    for line in lines:
        piece = line.strip()
        if not piece:
            continue
        if out and re.search(r"\w-$", out):
            out += piece
        elif out:
            out += " " + piece
        else:
            out = piece
    return norm_ws(out)


def split_prefs(value, same_rank_slash):
    """Turn a preference string into [(stallion, tag)] lists.  In the master file a '/' inside a preference
    line means alternatives at the same rank; in the index ' / ', ' or ' and '/' separate the 1st and 2nd
    preference (the caller decides via same_rank_slash)."""
    parts = re.split(r"\s*/\s*|\s+or\s+", value) if not same_rank_slash else re.split(r"\s*/\s*", value)
    out = []
    for part in parts:
        part = part.strip()
        if not part:
            continue
        tag = None
        m = re.search(r"\(([^)]*)\)", part)
        if m:
            tag = norm_ws(m.group(1))
            part = part[:m.start()] + part[m.end():]
        name = horse_name(part)
        if name:
            out.append((name, tag))
    return out


# ----------------------------------------------------------------------------------------------------
# master file
# ----------------------------------------------------------------------------------------------------
PAGE_RE = re.compile(r"^===== PAGE (\d+) \(chars=\d+\) =====[ \t]*$", re.M)
MASTER_TITLE = "WATHNAN RACING MATING SUGGESTIONS"
HEADER_RE = re.compile(r"^(?P<name>[A-Z][A-Z'’ ]+?)\s*\((?:(?P<yob>[^()]*?\d{4})\s+)?(?P<sire>[^()]+?)"
                       r"\s+(?:x|ex)\s+(?P<dam>[^()]+?)\s+by\s+(?P<damsire>[^()]+?)\)\s*$")
STATUS_RE = re.compile(r"^(?:Status\s*:\s*)?i/f\s+(?:to\s+)?(?P<sire>.+?)\s+lsd\s+(?P<d>\d{1,2})/(?P<m>\d{1,2})"
                       r"/(?P<y>\d{2,4})\s*$", re.I)
MAIDEN_RE = re.compile(r"^(?:Status\s*:\s*)?MAIDEN\s*$", re.I)
BARREN_RE = re.compile(r"^(?:Status\s*:\s*)?BARREN(?:\s*\((?:to\s+)?(?P<sire>[^)]+)\))?\s*$", re.I)
PREF_RE = re.compile(r"^(?P<rank>\d)(?:st|nd|rd|th)\s+preference(?:\s+(?P<region>[A-Za-z]+))?\s*:\s*(?P<value>.*?)\s*$",
                     re.I)
LABEL_RE = re.compile(r"^(?P<label>Race Record|Pedigree|Produce Record|Analysis)\s*:\s*(?P<rest>.*)$")
MIDLINE_LABEL_RE = re.compile(r"(?<!^)\b(Race Record|Pedigree|Produce Record|Analysis)\s*:")
CAPTION_RE = re.compile(r"^[A-Z][A-Za-z'’]*(?: [A-Za-z'’]+){0,4}$")
SECTION_KEYS = {"Race Record": "raceRecord", "Pedigree": "pedigree", "Produce Record": "produceRecord",
                "Analysis": "analysis"}


def load_master_pages(path):
    txt = fix_ligatures(open(path, encoding="utf-8").read())
    parts = PAGE_RE.split(txt)
    return [(int(parts[i]), parts[i + 1]) for i in range(1, len(parts), 2)]


def split_captions(text):
    """Separate a page into prose lines and the trailing stallion-name captions of the pasted 5x5 grids.
    A caption block is one or more short name-only lines at the very end of the page, preceded by a blank
    line (or forming the whole page)."""
    lines = [l.rstrip() for l in text.split("\n")]
    while lines and not lines[-1].strip():
        lines.pop()
    i = len(lines)
    while i > 0 and len(lines[i - 1].strip()) <= 30 and CAPTION_RE.match(lines[i - 1].strip()):
        i -= 1
    if i < len(lines) and (i == 0 or not lines[i - 1].strip()):
        return lines[:i], [l.strip() for l in lines[i:]]
    return lines, []


def parse_master(path):
    """-> {mare name: report dict}"""
    reports, cur = [], None
    for pno, text in load_master_pages(path):
        lines, captions = split_captions(text)
        for line in lines:
            s = line.strip()
            if s == MASTER_TITLE:
                continue
            m = HEADER_RE.match(s)
            if m:
                cur = dict(header=m, name=horse_name(m["name"]), lines=[], captions=[], pages=[pno])
                reports.append(cur)
                continue
            if cur is not None:
                cur["lines"].append(line)
        if cur is not None:
            cur["captions"].extend(captions)
            if pno not in cur["pages"]:
                cur["pages"].append(pno)
    out = {}
    for r in reports:
        rep = parse_master_report(r)
        if rep["name"] in out:
            rep["warnings"].append("duplicate report for this mare in the master file")
        out[rep["name"]] = rep
    return out


def parse_master_status(s):
    m = STATUS_RE.match(s)
    if m:
        return dict(status="in_foal", coveringSire=horse_name(m["sire"]),
                    lastServiceDate=iso_date(f"{m['d']}/{m['m']}/{m['y']}"))
    if MAIDEN_RE.match(s):
        return dict(status="maiden", coveringSire=None, lastServiceDate=None)
    m = BARREN_RE.match(s)
    if m:
        return dict(status="barren", coveringSire=horse_name(m["sire"]) if m["sire"] else None,
                    lastServiceDate=None)
    return None


def parse_master_report(r):
    h, warnings = r["header"], []
    name = r["name"]
    yob = None
    if h["yob"]:
        ym = re.search(r"(?:19|20)\d{2}", h["yob"])
        yob = int(ym.group(0)) if ym else None
    lines = r["lines"]
    i = 0
    pre = []
    while i < len(lines) and not LABEL_RE.match(lines[i].strip()):
        pre.append(lines[i])
        i += 1
    body = lines[i:]

    status = status_line = None
    prefs, physical = [], []
    for line in pre:
        s = norm_ws(line)
        if not s:
            continue
        pm = PREF_RE.match(s)
        if pm:
            region = REGION_ALIASES.get(pm["region"].upper()) if pm["region"] else None
            if pm["region"] and not region:
                warnings.append(f"unknown preference region {pm['region']!r}")
            for stallion, tag in split_prefs(pm["value"], same_rank_slash=True):
                prefs.append(dict(rank=int(pm["rank"]), stallion=stallion, region=region, tag=tag))
            continue
        st = parse_master_status(s)
        if st and status is None:
            status, status_line = st, re.sub(r"^Status\s*:\s*", "", s)
            continue
        if horse_name(s) == name:
            continue  # the mare's own name, printed as the caption of her photograph
        physical.append(s)
    if status is None:
        warnings.append("no status line found")
    if not prefs:
        warnings.append("no preference lines found")

    sections, cur = {}, None
    for line in body:
        lm = LABEL_RE.match(line.strip())
        if lm:
            cur = lm["label"]
            if cur in sections:
                warnings.append(f"duplicate '{cur}:' section")
            sections.setdefault(cur, [])
            if lm["rest"].strip():
                sections[cur].append(lm["rest"])
            continue
        if MIDLINE_LABEL_RE.search(line):
            warnings.append(f"section label inside a line: {norm_ws(line)[:60]!r}")
        if cur is None:
            if line.strip():
                warnings.append(f"text before the first section label: {norm_ws(line)[:60]!r}")
            continue
        sections[cur].append(line)
    text = {SECTION_KEYS[k]: join_wrapped_lines(v) for k, v in sections.items()}
    for key in ("Race Record", "Pedigree", "Analysis"):
        if key not in sections:
            warnings.append(f"no '{key}:' section")

    grid, seen = [], set()
    for c in r["captions"]:
        n = horse_name(c)
        if n not in seen:
            seen.add(n)
            grid.append(n)
    return dict(name=name, yob=yob, sire=horse_name(h["sire"]), dam=horse_name(h["dam"]),
                damsire=horse_name(h["damsire"]), pages=r["pages"], status=status, statusLine=status_line,
                preferences=prefs, physical=physical, gridStallions=grid,
                raceRecord=text.get("raceRecord"), pedigree=text.get("pedigree"),
                produceRecord=text.get("produceRecord"), analysis=text.get("analysis"), warnings=warnings)


# ----------------------------------------------------------------------------------------------------
# index + summary workbook
# ----------------------------------------------------------------------------------------------------
def parse_index():
    rows = {}
    for line in INDEX_TEXT.strip().splitlines():
        cells = [c.strip() for c in line.split("|")]
        if len(cells) != 11:
            raise ValueError(f"index row has {len(cells)} cells: {line!r}")
        raw_name, yob, sire, dam, location, country, foaled, plan, status, lsd, suggestion = cells
        m = re.match(r"^(.*?)(?:\s*\(([A-Z]+)\))?$", raw_name)
        name = horse_name(m.group(1))
        name = INDEX_NAME_FIXES.get(name, name)
        rows[name] = dict(name=name, indexName=raw_name, foaledIn=m.group(2), yob=int(yob), sire=horse_name(sire),
                          dam=horse_name(dam), location=None if blank(location) else location,
                          standsIn=None if blank(country) else STANDS_IN.get(country.upper(), country.upper()),
                          foaled2026=None if blank(foaled) or foaled.lower() == "maiden" else foaled,
                          plan2026=None if blank(plan) else horse_name(plan),
                          status=None if blank(status) else status,
                          lsd=None if blank(lsd) else lsd, suggestion=suggestion,
                          maiden="maiden" in (foaled + plan + status).lower())
    return rows


def load_workbook(path):
    import openpyxl  # PYTHONPATH=<scratch>/pylib
    wb = openpyxl.load_workbook(path, data_only=True)
    grid, fillies = {}, []
    ws = wb["Mares"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or blank(row[0]):
            continue
        name = horse_name(str(row[0]))
        grid[name] = dict(sire=horse_name(str(row[1] or "")), foaled=norm_ws(str(row[2])) if not blank(row[2]) else None,
                          inFoalTo=norm_ws(str(row[3])) if not blank(row[3]) else None,
                          stallions=[horse_name(str(c)) for c in row[4:7] if not blank(c)])
    ws = wb["Fillies in training"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or blank(row[0]):
            continue
        fillies.append(dict(name=horse_name(str(row[0])), age=int(row[1]) if not blank(row[1]) else None,
                            sire=horse_name(str(row[2] or "")), dam=horse_name(str(row[3] or "")),
                            damsire=horse_name(str(row[4] or "")), notes=norm_ws(str(row[5])) if not blank(row[5]) else ""))
    return grid, fillies


GRID_FOAL_RE = re.compile(r"^(?P<sire>[A-Za-z'’ ]+?)\s+(?:(?P<sex>[fc])\s+)?(?P<d>\d{1,2})\.(?P<m>\d{1,2})$")


def parse_grid_foal(s):
    """'Frankel f 07.02' -> (FRANKEL, 'filly', '2026-02-07'); 'Justify 08.02' -> (JUSTIFY, None, '2026-02-08')."""
    if not s:
        return None, None, None
    m = GRID_FOAL_RE.match(s.strip())
    if not m:
        return None, None, None
    sex = {"f": "filly", "c": "colt"}.get(m["sex"]) if m["sex"] else None
    return horse_name(m["sire"]), sex, iso_date(f"{m['d']}.{m['m']}.{PREV_SEASON}")


def foal_sex_from_text(sire, texts):
    """Find 'a Frankel filly', 'filly foal by Isaac Shelby', 'an ISAAC SHELBY filly foal' in the prose."""
    if not sire:
        return None
    s = re.escape(sire.title()).replace("\\ ", r"\s+")
    pats = [rf"\b(filly|colt)\s+foal\s+by\s+{s}\b", rf"\b{s}\s+(filly|colt)\b", rf"\bfoaled\b[^.]*?\b{s}\s+(filly|colt)\b"]
    for t in texts:
        if not t:
            continue
        for p in pats:
            m = re.search(p, t, re.I)
            if m:
                return m.group(1).lower()
    return None


# ----------------------------------------------------------------------------------------------------
# docx pack (zip on Google Drive, fetched by range request; central directory in ziptail.bin)
# ----------------------------------------------------------------------------------------------------
def zip_entries(scratch):
    tail = os.path.join(scratch, "ziptail.bin")
    if not os.path.exists(tail):
        start = ZIP_SIZE - ZIP_TAIL_BYTES
        raw = subprocess.run(["curl", "-sS", "-L", "-r", f"{start}-{ZIP_SIZE - 1}", "--max-time", "180", ZIP_URL],
                             capture_output=True, check=True).stdout
        if len(raw) != ZIP_TAIL_BYTES:
            raise RuntimeError(f"zip tail fetch returned {len(raw)} bytes")
        open(tail, "wb").write(raw)
    b = open(tail, "rb").read()
    tail_start = ZIP_SIZE - len(b)
    e = b.rfind(b"PK\x05\x06")
    (_, _, _, n_total, cd_size, cd_offset, _) = struct.unpack("<HHHHIIH", b[e + 4:e + 22])
    j = cd_offset - tail_start
    entries = []
    for _ in range(n_total):
        (ver, vn, flag, method, mt, md, crc, csize, usize, nlen, elen, cl, dsk, ia, ea, lho) = \
            struct.unpack("<HHHHHHIIIHHHHHII", b[j + 4:j + 46])
        name = b[j + 46:j + 46 + nlen].decode("utf-8", "ignore")
        entries.append(dict(name=name, method=method, csize=csize, usize=usize, lho=lho))
        j += 46 + nlen + elen + cl
    return entries


def fetch_zip_entry(ent):
    start = ent["lho"]
    end = start + 30 + 2048 + ent["csize"]
    raw = subprocess.run(["curl", "-sS", "-L", "-r", f"{start}-{end}", "--max-time", "180", ZIP_URL],
                         capture_output=True).stdout
    if raw[:4] != b"PK\x03\x04":
        raise RuntimeError(f"bad local header for {ent['name']}: {raw[:120]!r}")
    nlen, elen = struct.unpack("<HH", raw[26:30])
    data = raw[30 + nlen + elen:30 + nlen + elen + ent["csize"]]
    if ent["method"] == 8:
        import zlib
        data = zlib.decompress(data, -15)
    if len(data) != ent["usize"]:
        raise RuntimeError(f"size mismatch for {ent['name']}: {len(data)} != {ent['usize']}")
    return data


def docx_files(scratch, allow_fetch, log):
    """-> {mare folder name (UPPER): local path of '<MARE> 27.docx'}"""
    cache = os.path.join(scratch, "fixture", "docx")
    os.makedirs(cache, exist_ok=True)
    out, fetched = {}, 0
    for ent in zip_entries(scratch):
        m = re.match(re.escape(ZIP_PACK_PREFIX) + r"([^/]+)/([^/]+ 27\.docx)$", ent["name"])
        if not m:
            continue
        mare, base = horse_name(m.group(1)), m.group(2)
        local = None
        for d in (cache, os.path.join(scratch, "zip")):
            p = os.path.join(d, base)
            if os.path.exists(p) and os.path.getsize(p) == ent["usize"]:
                local = p
                break
        if local is None and allow_fetch:
            if fetched:
                time.sleep(1.0)
            log(f"fetching {ent['name']} ({ent['usize']} bytes)")
            data = fetch_zip_entry(ent)
            local = os.path.join(cache, base)
            open(local, "wb").write(data)
            fetched += 1
        if local is None:
            log(f"WARNING: {ent['name']} not cached and fetching disabled")
            continue
        out[mare] = local
    return out


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _para_text(p):
    out = []
    for el in p.iter():
        if el.tag == W + "t":
            out.append(el.text or "")
        elif el.tag == W + "tab":
            out.append("\t")
        elif el.tag in (W + "br", W + "cr"):
            out.append("\n")
        elif el.tag == W + "noBreakHyphen":
            out.append("-")
    return "".join(out).replace("﻿", "").replace("\xa0", " ")


def docx_blocks(path):
    """Top-level body blocks in order: ('p', text) or ('tbl', [[cell, ...], ...])."""
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    blocks = []
    for child in root.find(W + "body"):
        if child.tag == W + "p":
            blocks.append(("p", _para_text(child)))
        elif child.tag == W + "tbl":
            rows = []
            for tr in child.iter(W + "tr"):
                rows.append(["\n".join(_para_text(p) for p in tc.iter(W + "p")) for tc in tr.findall(W + "tc")])
            blocks.append(("tbl", rows))
    return blocks


# "CIRCIOS 2022 by Kingman ex Could It Be Love (War Front)"; tolerates "SOPRANO (2021) by ..." and a doubled ")".
DOCX_TITLE_RE = re.compile(r"^(?P<name>[A-Z][A-Z'’ ]+?)\s+\(?(?P<yob>\d{4})\)?\s+by\s+(?P<sire>.+?)\s+ex\s+"
                           r"(?P<dam>.+?)\s*\((?P<damsire>[^()]+)\)[)\s]*$")
STALLION_HDR_RE = re.compile(r"^(?P<rank>\d{1,2})[.)]?\s+(?P<stallion>[A-Z][A-Z'’ ]+?)\s+(?P<header>by\s+.+?\s+ex\s+.+?"
                             r"\([^()]+\))(?P<rest>.*)$", re.S)
FEE_RE = re.compile(r"(?P<year>\d{4})\s*fee\s*:\s*(?P<fee>[^\n\t]*)", re.I)
# A statistic line (as opposed to the argument paragraph) starts with a number ("17.7% SW/R", "40 G1Ws from ...",
# "3x4 Green Desert"), with a "<Sire> / <Line>" cross, with "Direct cross", or with "<Sire> [with <Line>] more
# generally / generally / unproven / direct cross".  Patterns are anchored at the start because argument prose
# routinely contains "the direct cross" and "more generally".
STAT_RE = re.compile(r"""
      ^\d
    | ^[A-Z][A-Za-z'’.;]*(?:\s+[A-Za-z'’.]+){0,4}\s*/\s*[A-Z]
    | ^Direct\s+cross\b
    | ^(?:[A-Z][\w'’]*\s+){1,6}(?:more\s+gene\w*|generally|more\s+loosely|unproven|direct\s+cross)\b
    | ^(?:[A-Z][\w'’]*\s+){1,6}with\s+(?:[A-Z][\w'’]*\s+){1,4}(?:\d|more\s+gene\w*|generally|no\s+further|unproven|direct)
    | ^[A-Z][^.]{0,40}\bSW/Rs?\b
    """, re.X)
STAT_CONTINUES_RE = re.compile(r"(?:,|;|:|-|–|\band|\bgenerally|\bcross)$")
INBREEDING_RE = re.compile(r"^\d+\s*[xX]\s*\d+(?:\s*[xX]\s*\d+)?\s+\S")
PROFILE_HEADING_RE = re.compile(r"^profile\s*:?$", re.I)
# Profile paragraph classification: a paragraph about the mare's own foals/yearlings ("Her first is a Wootton
# Bassett yearling filly ..."), and a closing paragraph describing her physique, recognised at sentence level
# (a physical adjective in a sentence about "she"/"her"/the mare/filly, ignoring sentences that start with "I").
PRODUCE_RE = re.compile(r"^(?:Her first (?:foal|is)\b|Her (?:first|second|third|fourth) foal\b|She is the dam of\b|"
                        r"Her (?:only|two|three) foals?\b|Her (?:yearling|foal|2yo|two[- ]year[- ]old)\b)", re.I)
PHYS_KW_RE = re.compile(r"\b(?:not very big|big|small|compact|rangy|rangey|leggy|legs?|short of leg|light[- ]framed|heavy|tall|"
                        r"scope|scopey|good[- ]looking|attractive|correct|straightforward|walks?|bone|\d{2}(?:\.\d)? ?hh|hands|"
                        r"neat|plain|size|height|substance|outlook|knees|offset|physically|physical|conformation|frame|athletic)\b",
                        re.I)
PHYS_SUBJECT_RE = re.compile(r"\b(?:she|her|mare|filly|type)\b", re.I)


def is_physical_para(p, mare_name):
    for sent in re.split(r"(?<=[.!?])\s+", p):
        if re.match(r"^I\b", sent):
            continue
        if PHYS_KW_RE.search(sent) and (PHYS_SUBJECT_RE.search(sent) or (mare_name and mare_name.lower() in sent.lower())):
            return True
    return False


def inbreeding_tag(stats):
    """'3x5 inbreeding to Danehill - Frankel with Danehill 30 SW ...' -> '3x5 inbreeding to Danehill'."""
    for x in stats:
        if INBREEDING_RE.match(x):
            return re.split(r"\s+[-–]\s+|(?<=[a-z])\.\s|;", x)[0].strip().rstrip(".")
    return None


def parse_docx(path, expected_name=None):
    warnings = []
    blocks = docx_blocks(path)
    paras = [(k, v) for k, v in blocks]
    # title
    title = None
    ti = None
    for i, (k, v) in enumerate(paras):
        if k == "p" and norm_ws(v):
            title = norm_ws(v)
            ti = i
            break
    if title is None:
        return None, ["empty document"]
    tm = DOCX_TITLE_RE.match(title)
    head = {}
    if tm:
        head = dict(name=horse_name(tm["name"]), yob=int(tm["yob"]), sire=horse_name(tm["sire"]),
                    dam=horse_name(tm["dam"]), damsire=horse_name(tm["damsire"]))
        if expected_name and head["name"] != expected_name:
            warnings.append(f"title names {head['name']!r}, folder says {expected_name!r}")
    else:
        warnings.append(f"title line not in '<MARE> <YOB> by <Sire> ex <Dam> (<Damsire>)' form: {title!r}")

    # profile paragraphs: after the 'Profile' heading (or the title) up to the first stallion header
    profile, stallions, ruled_out = [], [], []
    cur = None          # current stallion block
    for k, v in paras[ti + 1:]:
        if k == "tbl":
            rows = [[norm_ws(c) for c in row] for row in v]
            rows = [r for r in rows if any(r)]
            if rows and re.search(r"ruled\s*out|other stallions", " ".join(rows[0]), re.I):
                for r in rows[1:]:
                    stallion = horse_name(r[0]) or None
                    reason = r[1] if len(r) > 1 else ""
                    if stallion is None:
                        warnings.append(f"ruled-out row without a stallion name: {reason!r}")
                    ruled_out.append(dict(stallion=stallion, reason=reason))
                if not rows[1:]:
                    warnings.append("ruled-out table has no rows")
            else:
                warnings.append(f"unexpected table ignored: {rows[:1]!r}")
            continue
        raw = v
        s = norm_ws(raw)
        if not s:
            continue
        if cur is None and not stallions and PROFILE_HEADING_RE.match(s):
            continue
        hm = STALLION_HDR_RE.match(s if "\n" not in raw else raw.strip())
        if hm:
            cur = dict(rank=int(hm["rank"]), stallion=horse_name(hm["stallion"]), header=norm_ws(hm["header"]),
                       fee=None, feeYear=None, stats=[], text=[], phase="stats")
            stallions.append(cur)
            rest = hm["rest"]
            fm = FEE_RE.search(rest)
            if fm:
                cur["fee"], cur["feeYear"] = norm_ws(fm["fee"]) or None, int(fm["year"])
                rest = rest[:fm.start()] + rest[fm.end():]
            pieces = [norm_ws(p) for p in re.split(r"[\n\t]", rest)]
            _absorb_stat_pieces(cur, [p for p in pieces if p], warnings)
            continue
        if cur is None:
            profile.append(s)          # profile paragraph
            continue
        if cur["phase"] == "stats":
            fm = FEE_RE.search(raw)
            if fm and cur["fee"] is None and len(s) < 60:
                cur["fee"], cur["feeYear"] = norm_ws(fm["fee"]) or None, int(fm["year"])
                continue
            pieces = [norm_ws(p) for p in re.split(r"[\n\t]", raw)]
            _absorb_stat_pieces(cur, [p for p in pieces if p], warnings)
        else:
            _append_para(cur["text"], s)
    if not stallions:
        warnings.append("no numbered stallion blocks found")
    for st in stallions:
        if not st["text"]:
            warnings.append(f"{st['stallion']}: no argument paragraph")
        if st["fee"] is None:
            warnings.append(f"{st['stallion']}: no fee")
        if not st["stats"]:
            warnings.append(f"{st['stallion']}: no statistic lines")
        for x in st["stats"]:
            if len(x) > 260:
                warnings.append(f"{st['stallion']}: very long statistic line ({len(x)} chars)")
        st.pop("phase", None)

    # merge run-on profile paragraphs (a paragraph starting in lower case continues the previous one)
    merged = []
    for p in profile:
        _append_para(merged, p)
    profile = merged
    race = profile[0] if profile else None
    rest = profile[1:]
    physical = None
    if rest and is_physical_para(rest[-1], head.get("name") or expected_name):
        physical = rest.pop()
    produce = [p for p in rest if PRODUCE_RE.match(p)]
    pedigree = [p for p in rest if not PRODUCE_RE.match(p)]
    if not profile:
        warnings.append("no profile paragraphs")
    return dict(title=title, head=head, raceRecord=race, pedigree=pedigree, produce=produce, physical=physical,
                stallions=stallions, ruledOut=ruled_out), warnings


def _append_para(paras, s):
    if paras and s[:1].islower():
        paras[-1] = paras[-1] + " " + s
    else:
        paras.append(s)


def _absorb_stat_pieces(cur, pieces, warnings):
    for p in pieces:
        if not re.search(r"[A-Za-z0-9]", p):
            continue  # stray punctuation such as the "``" left in two headers
        if cur["phase"] == "stats":
            prev = cur["stats"][-1] if cur["stats"] else None
            if prev is not None and (STAT_CONTINUES_RE.search(prev) or p[:1].islower()):
                cur["stats"][-1] = prev + " " + p           # continuation of the previous statistic line
            elif STAT_RE.search(p):
                cur["stats"].append(p)
            elif prev is not None and len(p) < 60 and not p.endswith("."):
                cur["stats"][-1] = prev + " " + p           # short unpunctuated fragment
            else:
                cur["phase"] = "text"
                cur["text"].append(p)
        else:
            _append_para(cur["text"], p)


# ----------------------------------------------------------------------------------------------------
# Timeform export (repo) - foaling country fallback
# ----------------------------------------------------------------------------------------------------
def timeform_lookup(wanted):
    """wanted: {name: (yob, sire)} -> {name: (FoalingCountry, FoalingYear)}; scans horses-<yob±1>.csv.gz."""
    found = {}
    years = sorted({y + d for (y, _) in wanted.values() if y for d in (0, -1, 1)})
    for y in years:
        path = os.path.join(TIMEFORM_DIR, f"horses-{y}.csv.gz")
        if not os.path.exists(path):
            continue
        with gzip.open(path, "rt", encoding="utf-8", newline="") as fh:
            for row in csv.DictReader(fh):
                n = row.get("HorseName", "").upper()
                if n in wanted and n not in found:
                    yob, sire = wanted[n]
                    if sire and horse_name(row.get("Sire", "")) != sire:
                        continue
                    found[n] = (row.get("FoalingCountry") or None, int(row["FoalingYear"]) if row.get("FoalingYear") else None)
    return found


# ----------------------------------------------------------------------------------------------------
# assembly
# ----------------------------------------------------------------------------------------------------
def build(scratch, allow_fetch, log):
    master = parse_master(os.path.join(scratch, "master_text.txt"))
    index = parse_index()
    grid, fillies = load_workbook(os.path.join(scratch, "zip", "Wathnan summary.xlsx"))
    docx_paths = docx_files(scratch, allow_fetch, log)

    names = sorted(set(index) | set(master) | set(DOCX_ONLY_MARES))
    for n in master:
        if n not in index:
            log(f"WARNING: master report for {n} has no index row")
    for n in docx_paths:
        if n not in names:
            log(f"WARNING: docx for {n} matches no mare")

    # docx parse
    docx = {}
    for n, p in docx_paths.items():
        parsed, warns = parse_docx(p, expected_name=n)
        docx[n] = (parsed, warns, os.path.basename(p))
        log(f"docx {os.path.basename(p)}: " + ("; ".join(warns) if warns else "ok"))

    # foaling country fallback from the Timeform export
    wanted = {}
    for n in names:
        ix = index.get(n)
        if ix and ix["foaledIn"]:
            continue
        yob = (master.get(n) or {}).get("yob") or (ix or {}).get("yob") or ((docx.get(n) or (None,))[0] or {}).get("head", {}).get("yob")
        sire = (master.get(n) or {}).get("sire") or (ix or {}).get("sire") or ((docx.get(n) or (None,))[0] or {}).get("head", {}).get("sire")
        wanted[n] = (yob, sire)
    timeform = timeform_lookup(wanted) if wanted else {}

    fillies_by_name = {f["name"]: f for f in fillies}
    mares = []
    for n in names:
        ix, ms = index.get(n), master.get(n)
        dx, dwarns, dbase = docx.get(n, (None, [], None))
        dhead = (dx or {}).get("head") or {}
        mare_warnings = []

        yob = ms["yob"] if ms and ms["yob"] else None
        if yob is None and ix:
            yob = ix["yob"]
        if yob is None and dhead.get("yob"):
            yob = dhead["yob"]
        for src, val in (("index", (ix or {}).get("yob")), ("docx", dhead.get("yob"))):
            if val and yob and val != yob:
                mare_warnings.append(f"yob: {src} says {val}, using {yob}")
        sire = (ms or {}).get("sire") or (ix or {}).get("sire") or dhead.get("sire")
        dam = (ms or {}).get("dam") or (ix or {}).get("dam") or dhead.get("dam")
        damsire = (ms or {}).get("damsire") or dhead.get("damsire")
        for src, d in (("index", ix or {}), ("docx", dhead)):
            for key, val in (("sire", sire), ("dam", dam), ("damsire", damsire)):
                if d.get(key) and val and d[key] != val:
                    mare_warnings.append(f"{key}: {src} says {d[key]!r}, using {val!r}")
        foaled_in = (ix or {}).get("foaledIn")
        if not foaled_in and n in timeform:
            foaled_in = timeform[n][0]
            mare_warnings.append(f"foaledIn {foaled_in!r} taken from the Timeform export (index gives no country)")
        elif not foaled_in:
            mare_warnings.append("foaledIn unknown (no index suffix, not found in the Timeform export)")

        # 2026 season
        season = dict(status=None, coveringSire=None, lastServiceDate=None, foaledDate=None, foalSex=None, foalBy=None)
        if ix and not ix["maiden"]:
            season["status"] = {"in foal": "in_foal", "not in foal": "not_in_foal"}.get((ix["status"] or "").lower())
            if season["status"] is None:
                mare_warnings.append(f"index status {ix['status']!r} not understood")
            season["coveringSire"] = ix["plan2026"]
            season["lastServiceDate"] = iso_date(ix["lsd"])
            season["foaledDate"] = iso_date(ix["foaled2026"])
        else:
            season["status"] = "maiden"
        if ms and ms["status"]:
            mst = ms["status"]
            if mst["status"] == "barren":
                season["status"] = "barren"
            elif mst["status"] != season["status"] and not (mst["status"] == "in_foal" and season["status"] == "in_foal"):
                mare_warnings.append(f"status: master says {mst['status']!r}, index says {season['status']!r}")
            if mst["coveringSire"] and season["coveringSire"] and mst["coveringSire"] != season["coveringSire"]:
                mare_warnings.append(f"covering sire: master says {mst['coveringSire']!r}, index says {season['coveringSire']!r}")
            season["coveringSire"] = season["coveringSire"] or mst["coveringSire"]
            if mst["lastServiceDate"] and season["lastServiceDate"] and mst["lastServiceDate"] != season["lastServiceDate"]:
                mare_warnings.append(f"lsd: master says {mst['lastServiceDate']}, index says {season['lastServiceDate']}")
            season["lastServiceDate"] = season["lastServiceDate"] or mst["lastServiceDate"]
        g = grid.get(n)
        if g:
            foal_by, foal_sex, foal_date = parse_grid_foal(g["foaled"])
            if foal_by and season["foaledDate"] is None and ix and not ix["maiden"]:
                mare_warnings.append(f"summary grid records a 2026 foal ({g['foaled']!r}) but the index has no foaling date")
            if foal_by:
                season["foalBy"] = foal_by
                if foal_date and season["foaledDate"] and foal_date != season["foaledDate"]:
                    mare_warnings.append(f"foaling date: grid says {foal_date}, index says {season['foaledDate']} (index kept)")
                season["foalSex"] = foal_sex or foal_sex_from_text(
                    foal_by, [(ms or {}).get("produceRecord"), (ms or {}).get("raceRecord"), (ms or {}).get("analysis"),
                              (dx or {}).get("raceRecord")] + list((dx or {}).get("produce") or []) + list((dx or {}).get("pedigree") or []))
        elif ix and not ix["maiden"]:
            mare_warnings.append("not in the summary grid")

        # physical notes from every source
        physical = list((ms or {}).get("physical") or [])
        if dx and dx.get("physical"):
            physical.append(dx["physical"])
        fit = fillies_by_name.get(n)
        if fit and fit["notes"]:
            physical.append(fit["notes"])

        plans = []
        if ms:
            prefs = ms["preferences"]
            if not prefs and ix:
                prefs = [dict(rank=i + 1, stallion=s, region=None, tag=t)
                         for i, (s, t) in enumerate(split_prefs(ix["suggestion"], same_rank_slash=False))]
                ms["warnings"].append("preferences taken from the index (master file has none)")
            plans.append(dict(source="master", statusLine=ms["statusLine"], preferences=prefs,
                              gridStallions=ms["gridStallions"], raceRecord=ms["raceRecord"], pedigree=ms["pedigree"],
                              produceRecord=ms["produceRecord"], analysis=ms["analysis"], ruledOut=[], stallionNotes=[],
                              parseWarnings=ms["warnings"] + mare_warnings))
        elif ix:
            plans.append(dict(source="index", statusLine=None,
                              preferences=[dict(rank=i + 1, stallion=s, region=None, tag=t)
                                           for i, (s, t) in enumerate(split_prefs(ix["suggestion"], same_rank_slash=False))],
                              gridStallions=[], raceRecord=None, pedigree=None, produceRecord=None, analysis=None,
                              ruledOut=[], stallionNotes=[], parseWarnings=["no master file report"] + mare_warnings))
        if dx:
            prefs = []
            notes = []
            for st in dx["stallions"]:
                tag = inbreeding_tag(st["stats"])
                prefs.append(dict(rank=st["rank"], stallion=st["stallion"], region=None, tag=tag))
                notes.append(dict(stallion=st["stallion"], header=st["header"], fee=st["fee"], feeYear=st["feeYear"],
                                  stats=st["stats"], text="\n\n".join(st["text"]) if st["text"] else None))
            analysis = "\n\n".join(f"{st['stallion']} — " + "\n\n".join(st["text"]) for st in dx["stallions"] if st["text"])
            dw = list(dwarns) + ([] if ms else mare_warnings)
            plans.append(dict(source="docx", file=dbase, title=dx["title"], preferences=prefs,
                              raceRecord=dx["raceRecord"], pedigree="\n\n".join(dx["pedigree"]) or None,
                              produceRecord="\n\n".join(dx["produce"]) or None, analysis=analysis or None,
                              stallionNotes=notes, ruledOut=dx["ruledOut"], parseWarnings=dw))
        if not plans:
            log(f"WARNING: {n} has no plan at all")

        mares.append(dict(name=n, foaledIn=foaled_in, yob=yob, sire=sire, dam=dam, damsire=damsire,
                          location=(ix or {}).get("location"), standsIn=(ix or {}).get("standsIn"), role="broodmare",
                          physical=" ".join(physical), seasons={PREV_SEASON: season}, plans2027=plans))

    fillies_out = [dict(name=f["name"], age=f["age"], sire=f["sire"], dam=f["dam"], damsire=f["damsire"], notes=f["notes"])
                   for f in fillies if f["name"] not in set(names)]

    return dict(
        client="Wathnan Racing", userId="richardbrown1", season=SEASON,
        sources=dict(index="Matings suggestions index for 2027 (PDF)",
                     master="Wathnan Matings Master File for 2027 Season (PDF)",
                     workingFolder="OneDrive WATHNAN MATINGS 2026 (docx pack)",
                     summary="Wathnan summary.xlsx (summary grid + fillies in training)",
                     timeform="breeding/data/horses-<yob>.csv.gz (foaling country where the index gives none)",
                     builtAt=dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")),
        mares=mares, filliesInTraining=fillies_out)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--scratch", default=DEFAULT_SCRATCH, help="scratch dir holding master_text.txt, ziptail.bin, zip/, fixture/")
    ap.add_argument("--out", default=DEFAULT_OUT)
    ap.add_argument("--no-fetch", action="store_true", help="never hit Google Drive; use only cached docx")
    args = ap.parse_args()

    def log(msg):
        print(msg, file=sys.stderr, flush=True)

    fixture = build(args.scratch, not args.no_fetch, log)
    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    with open(args.out, "w", encoding="utf-8") as fh:
        json.dump(fixture, fh, ensure_ascii=False, indent=2)
        fh.write("\n")
    n_docx = sum(1 for m in fixture["mares"] for p in m["plans2027"] if p["source"] == "docx")
    n_master = sum(1 for m in fixture["mares"] for p in m["plans2027"] if p["source"] == "master")
    log(f"wrote {args.out}: {len(fixture['mares'])} mares, {n_master} master plans, {n_docx} docx plans, "
        f"{len(fixture['filliesInTraining'])} fillies in training")


if __name__ == "__main__":
    main()
